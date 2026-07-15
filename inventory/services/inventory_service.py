"""
Inventory service module containing business logic for inventory operations.

This module encapsulates the business logic for inventory operations,
separating it from views and models.
"""

from datetime import datetime
from typing import Dict, List, Optional, Tuple, Any, Union, Set

from django.db import transaction
from django.db.models import Q, Sum, Count
from django.utils import timezone
from django.core.files import File
from django.core.files.temp import NamedTemporaryFile
from django.forms import ValidationError
from django.contrib import messages
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
import logging
import csv
import re

from ..models import (
    InventoryItem, InventorySeries, Location, Manufacturer, Organization,
    Category, InventoryImport, Inspection, InspectionImport, AuditLog
)

logger = logging.getLogger('django')
WS_NAME = _('Inventory')
IGNORED_PREFIXES = [_('Test'), _('Sample')]
REQUIRED_INSPECTION_FIELDS: Set[str] = {"inspection_number", "inspection_date"}


class InventoryService:
    """
    Service class for inventory operations.
    
    Encapsulates business logic for managing inventory items,
    import/export operations, and inspection management.
    """
    
    @staticmethod
    def get_inventory_items_with_availability() -> List[Dict]:
        """
        Get all inventory items with their availability information.
        
        Returns:
            List[Dict]: List of inventory items with availability data
        """
        items = InventoryItem.objects.select_related(
            'manufacturer', 'category', 'location', 'owner'
        ).all()
        
        result = []
        for item in items:
            result.append({
                'id': item.id,
                'inventory_number': item.inventory_number,
                'description': item.description,
                'serial_number': item.serial_number,
                'manufacturer': item.manufacturer.name if item.manufacturer else None,
                'category': item.category.name if item.category else None,
                'location': item.location.full_path if item.location else None,
                'quantity': item.quantity,
                'status': item.status,
                'owner': item.owner.name if item.owner else None,
                'available_for_rent': item.available_for_rent,
                'reserved_quantity': item.reserved_quantity,
                'rented_quantity': item.rented_quantity,
                'available_quantity': (item.quantity or 0) - (item.reserved_quantity or 0) - (item.rented_quantity or 0)
            })
        
        return result
    
    @staticmethod
    def get_items_by_ids(item_ids: List[int]) -> List[InventoryItem]:
        """
        Get inventory items by their IDs.
        
        Args:
            item_ids: List of inventory item IDs
            
        Returns:
            List[InventoryItem]: List of inventory items
        """
        return InventoryItem.objects.filter(id__in=item_ids).select_related(
            'manufacturer', 'category', 'location', 'owner'
        )
    
    @staticmethod
    def get_inventory_item_by_number(inventory_number: str) -> Optional[InventoryItem]:
        """
        Get an inventory item by its inventory number.
        
        Args:
            inventory_number: The inventory number to search for
            
        Returns:
            InventoryItem or None: The found inventory item or None
        """
        try:
            return InventoryItem.objects.get(inventory_number=inventory_number)
        except InventoryItem.DoesNotExist:
            return None

    # Trailing digits are the running number, everything before them is the
    # series prefix. Used for numbers whose series is not configured, so that
    # "TEST-EMAIL-001" or "OK-000456/1" still stay on their own sequence.
    _NUMBER_PATTERN = re.compile(r'^(?P<prefix>.*?)(?P<digits>\d+)$')

    @staticmethod
    def match_series(inventory_number: str) -> Optional[InventorySeries]:
        """
        Return the active series an inventory number belongs to.

        Args:
            inventory_number: Number to look up

        Returns:
            InventorySeries or None: Longest matching active series, if any
        """
        number = inventory_number or ''
        best = None
        for series in InventorySeries.objects.filter(active=True):
            if not number.startswith(series.prefix):
                continue
            # Longest prefix wins, so "INV-SUB-" beats "INV-" for its numbers.
            if best is None or len(series.prefix) > len(best.prefix):
                best = series
        return best

    @staticmethod
    def is_valid_inventory_number(inventory_number: str) -> bool:
        """
        Check that a number belongs to a configured series and has digits.

        Args:
            inventory_number: Number to validate

        Returns:
            bool: True if valid, False otherwise
        """
        number = inventory_number or ''
        series = InventoryService.match_series(number)
        if series is None:
            return False
        return bool(re.match(r'^\d+', number[len(series.prefix):]))

    @staticmethod
    def generate_next_inventory_number(source_number: str) -> str:
        """
        Return the next free inventory number in the source number's series.

        Args:
            source_number: Number to derive the series and zero-padding from

        Returns:
            str: An inventory number that is not taken yet
        """
        next_value = None
        series = InventoryService.match_series(source_number)
        if series is not None:
            prefix = series.prefix
            width = series.padding
        else:
            # Series is not configured - fall back to the number's own shape.
            match = InventoryService._NUMBER_PATTERN.match(source_number or '')
            if match:
                prefix = match.group('prefix')
                width = len(match.group('digits'))
            else:
                # Nothing to increment - open a sub-sequence under the number.
                prefix = f'{source_number}-'
                width = 1
                next_value = 1

        taken = set(
            InventoryItem.objects
            .filter(inventory_number__startswith=prefix)
            .values_list('inventory_number', flat=True)
        )

        if next_value is None:
            highest = 0
            for number in taken:
                # startswith also matches longer prefixes (e.g. "OK-000456/1"
                # for prefix "OK-"), so only count the same series.
                other = InventoryService._NUMBER_PATTERN.match(number)
                if other and other.group('prefix') == prefix:
                    highest = max(highest, int(other.group('digits')))
            next_value = highest + 1

        while f'{prefix}{next_value:0{width}d}' in taken:
            next_value += 1
        return f'{prefix}{next_value:0{width}d}'

    @staticmethod
    def create_inventory_item(
        inventory_number: str,
        description: str = '',
        serial_number: str = '',
        manufacturer: Optional[Manufacturer] = None,
        category: Optional[Category] = None,
        location: Optional[Location] = None,
        quantity: int = 1,
        status: str = 'in_stock',
        owner: Optional[Organization] = None,
        inventory_number_owner: str = '',
        purchase_date: Optional[datetime] = None,
        purchase_cost: Optional[float] = None,
        available_for_rent: bool = False
    ) -> InventoryItem:
        """
        Create a new inventory item.
        
        Args:
            inventory_number: Unique inventory number
            description: Item description
            serial_number: Serial number
            manufacturer: Manufacturer object
            category: Category object
            location: Location object
            quantity: Quantity of the item
            status: Status of the item
            owner: Owner organization
            inventory_number_owner: Owner's inventory number
            purchase_date: Purchase date
            purchase_cost: Purchase cost
            available_for_rent: Whether the item is available for rent
            
        Returns:
            InventoryItem: The created inventory item
        """
        item = InventoryItem.objects.create(
            inventory_number=inventory_number,
            description=description,
            serial_number=serial_number,
            manufacturer=manufacturer,
            category=category,
            location=location,
            quantity=quantity,
            status=status,
            owner=owner,
            inventory_number_owner=inventory_number_owner,
            purchase_date=purchase_date,
            purchase_cost=purchase_cost,
            available_for_rent=available_for_rent
        )
        
        return item
    
    @staticmethod
    def update_inventory_item(
        item_id: int,
        **kwargs
    ) -> Optional[InventoryItem]:
        """
        Update an existing inventory item.
        
        Args:
            item_id: ID of the item to update
            **kwargs: Fields to update
            
        Returns:
            InventoryItem or None: The updated item or None if not found
        """
        try:
            item = InventoryItem.objects.get(id=item_id)
            
            for field, value in kwargs.items():
                if hasattr(item, field):
                    setattr(item, field, value)
            
            item.save()
            return item
        except InventoryItem.DoesNotExist:
            return None
    
    @staticmethod
    def delete_inventory_item(item_id: int) -> bool:
        """
        Delete an inventory item.
        
        Args:
            item_id: ID of the item to delete
            
        Returns:
            bool: True if deleted successfully, False otherwise
        """
        try:
            item = InventoryItem.objects.get(id=item_id)
            item.delete()
            return True
        except InventoryItem.DoesNotExist:
            return False
    
    @staticmethod
    def get_or_create_manufacturer(name: str, description: str = '') -> Manufacturer:
        """
        Get or create a manufacturer.
        
        Args:
            name: Manufacturer name
            description: Manufacturer description
            
        Returns:
            Manufacturer: The manufacturer object
        """
        manufacturer, created = Manufacturer.objects.get_or_create(
            name=name,
            defaults={'description': description}
        )
        return manufacturer
    
    @staticmethod
    def get_or_create_organization(name: str, description: str = '') -> Organization:
        """
        Get or create an organization.
        
        Args:
            name: Organization name
            description: Organization description
            
        Returns:
            Organization: The organization object
        """
        organization, created = Organization.objects.get_or_create(
            name=name,
            defaults={'description': description}
        )
        return organization
    
    @staticmethod
    def get_or_create_category(name: str, description: str = '') -> Category:
        """
        Get or create a category.
        
        Args:
            name: Category name
            description: Category description
            
        Returns:
            Category: The category object
        """
        category, created = Category.objects.get_or_create(
            name=name,
            defaults={'description': description}
        )
        return category
    
    @staticmethod
    def get_or_create_location_by_path(path_str: str) -> Optional[Location]:
        """
        Get or create a location by hierarchical path.
        
        Args:
            path_str: Hierarchical path (e.g., "Building A -> Floor 1 -> Room 101")
            
        Returns:
            Location or None: The location object or None if path is empty
        """
        if not path_str:
            return None
        
        return Location.objects.get_or_create_by_path(path_str)
    
    @staticmethod
    def validate_inventory_file(file) -> None:
        """
        Validate an inventory import file.
        
        Args:
            file: The file to validate
            
        Raises:
            ValidationError: If the file is invalid
        """
        validate_inventory_import(file)
    
    @staticmethod
    def import_inventory_data(request, file, import_obj: InventoryImport) -> Dict[str, Union[int, str]]:
        """
        Import inventory data from a file.
        
        Args:
            request: Django request object
            file: The file to import
            import_obj: InventoryImport object for tracking
            
        Returns:
            Dict: Import results with created, skipped counts and error log
        """
        return inventory_import(request, file, import_obj)
    
    @staticmethod
    def validate_inspection_file(file) -> None:
        """
        Validate an inspection import file.
        
        Args:
            file: The file to validate
            
        Raises:
            ValidationError: If the file is invalid
        """
        validate_inspection_import(file)
    
    @staticmethod
    def import_inspection_data(
        request, 
        file, 
        import_obj: InspectionImport
    ) -> Dict[str, Union[int, str]]:
        """
        Import inspection data from a file.
        
        Args:
            request: Django request object
            file: The file to import
            import_obj: InspectionImport object for tracking
            
        Returns:
            Dict: Import results with created, skipped counts and error log
        """
        return inspection_import(request, file, import_obj)
    
    @staticmethod
    def create_inspection(
        inspection_number: str,
        inventory_item: Optional[InventoryItem] = None,
        manufacturer: str = '',
        device_type: str = '',
        room: str = '',
        target_part: str = 'device',
        inspection_date: Optional[datetime] = None,
        result: str = ''
    ) -> Inspection:
        """
        Create a new inspection record.
        
        Args:
            inspection_number: Unique inspection number
            inventory_item: Related inventory item
            manufacturer: Manufacturer name
            device_type: Type of device
            room: Room location
            target_part: Part of device inspected
            inspection_date: Date of inspection
            result: Inspection result
            
        Returns:
            Inspection: The created inspection record
        """
        inspection = Inspection.objects.create(
            inspection_number=inspection_number,
            inventory_item=inventory_item,
            manufacturer=manufacturer,
            device_type=device_type,
            room=room,
            target_part=target_part,
            inspection_date=inspection_date or timezone.now().date(),
            result=result
        )
        
        return inspection
    
    @staticmethod
    def get_inspections_for_inventory_item(
        inventory_item: InventoryItem
    ) -> List[Inspection]:
        """
        Get all inspections for a specific inventory item.
        
        Args:
            inventory_item: The inventory item to get inspections for
            
        Returns:
            List[Inspection]: List of inspection records
        """
        return Inspection.objects.filter(
            inventory_item=inventory_item
        ).order_by('-inspection_date')
    
    @staticmethod
    def get_latest_inspection_for_inventory_item(
        inventory_item: InventoryItem
    ) -> Optional[Inspection]:
        """
        Get the latest inspection for a specific inventory item.
        
        Args:
            inventory_item: The inventory item to get inspection for
            
        Returns:
            Inspection or None: The latest inspection or None
        """
        return Inspection.objects.filter(
            inventory_item=inventory_item
        ).order_by('-inspection_date').first()
    
    @staticmethod
    def get_audit_logs_for_inventory_item(
        inventory_item: InventoryItem
    ) -> List[AuditLog]:
        """
        Get all audit logs for a specific inventory item.
        
        Args:
            inventory_item: The inventory item to get audit logs for
            
        Returns:
            List[AuditLog]: List of audit log records
        """
        return AuditLog.objects.filter(
            model_name="InventoryItem",
            object_id=str(inventory_item.id)
        ).order_by('-timestamp')
    
    @staticmethod
    def search_inventory_items(
        query: str = '',
        manufacturer_id: Optional[int] = None,
        category_id: Optional[int] = None,
        location_id: Optional[int] = None,
        owner_id: Optional[int] = None,
        status: Optional[str] = None,
        available_for_rent: Optional[bool] = None
    ) -> List[InventoryItem]:
        """
        Search inventory items based on various criteria.
        
        Args:
            query: Search query for inventory number, description, or serial number
            manufacturer_id: Filter by manufacturer ID
            category_id: Filter by category ID
            location_id: Filter by location ID
            owner_id: Filter by owner ID
            status: Filter by status
            available_for_rent: Filter by availability for rent
            
        Returns:
            List[InventoryItem]: List of matching inventory items
        """
        items = InventoryItem.objects.select_related(
            'manufacturer', 'category', 'location', 'owner'
        )
        
        if query:
            items = items.filter(
                Q(inventory_number__icontains=query) |
                Q(description__icontains=query) |
                Q(serial_number__icontains=query)
            )
        
        if manufacturer_id:
            items = items.filter(manufacturer_id=manufacturer_id)
        
        if category_id:
            items = items.filter(category_id=category_id)
        
        if location_id:
            items = items.filter(location_id=location_id)
        
        if owner_id:
            items = items.filter(owner_id=owner_id)
        
        if status:
            items = items.filter(status=status)
        
        if available_for_rent is not None:
            items = items.filter(available_for_rent=available_for_rent)
        
        return items.all()
    
    @staticmethod
    def get_inventory_statistics() -> Dict[str, Any]:
        """
        Get overall inventory statistics.
        
        Returns:
            Dict: Dictionary containing various statistics
        """
        total_items = InventoryItem.objects.count()
        in_stock_items = InventoryItem.objects.filter(status='in_stock').count()
        rented_items = InventoryItem.objects.filter(status='rented').count()
        defect_items = InventoryItem.objects.filter(status='defect').count()
        written_off_items = InventoryItem.objects.filter(status='written_off').count()
        
        available_for_rent_items = InventoryItem.objects.filter(available_for_rent=True).count()
        
        total_quantity = InventoryItem.objects.aggregate(total=Sum('quantity'))['total'] or 0
        total_reserved = InventoryItem.objects.aggregate(total=Sum('reserved_quantity'))['total'] or 0
        total_rented = InventoryItem.objects.aggregate(total=Sum('rented_quantity'))['total'] or 0
        total_available = total_quantity - total_reserved - total_rented
        
        return {
            'total_items': total_items,
            'in_stock_items': in_stock_items,
            'rented_items': rented_items,
            'defect_items': defect_items,
            'written_off_items': written_off_items,
            'available_for_rent_items': available_for_rent_items,
            'total_quantity': total_quantity,
            'total_reserved': total_reserved,
            'total_rented': total_rented,
            'total_available': total_available,
            'items_by_status': {
                'in_stock': in_stock_items,
                'rented': rented_items,
                'defect': defect_items,
                'written_off': written_off_items,
            }
        }
    
    @staticmethod
    def update_inventory_quantities(
        inventory_item: InventoryItem,
        reserved_quantity: Optional[int] = None,
        rented_quantity: Optional[int] = None
    ) -> InventoryItem:
        """
        Update the reserved and rented quantities of an inventory item.
        
        Args:
            inventory_item: The inventory item to update
            reserved_quantity: New reserved quantity
            rented_quantity: New rented quantity
            
        Returns:
            InventoryItem: The updated inventory item
        """
        if reserved_quantity is not None:
            inventory_item.reserved_quantity = reserved_quantity
        
        if rented_quantity is not None:
            inventory_item.rented_quantity = rented_quantity
        
        inventory_item.save(update_fields=['reserved_quantity', 'rented_quantity'])
        return inventory_item
    
    @staticmethod
    def reserve_inventory_item(
        inventory_item: InventoryItem,
        quantity: int
    ) -> bool:
        """
        Reserve a quantity of an inventory item.
        
        Args:
            inventory_item: The inventory item to reserve
            quantity: Quantity to reserve
            
        Returns:
            bool: True if reservation was successful, False otherwise
        """
        available_quantity = (inventory_item.quantity or 0) - (inventory_item.reserved_quantity or 0) - (inventory_item.rented_quantity or 0)
        
        if available_quantity >= quantity:
            inventory_item.reserved_quantity = (inventory_item.reserved_quantity or 0) + quantity
            inventory_item.save(update_fields=['reserved_quantity'])
            return True
        
        return False
    
    @staticmethod
    def release_inventory_reservation(
        inventory_item: InventoryItem,
        quantity: int
    ) -> bool:
        """
        Release a reservation for an inventory item.
        
        Args:
            inventory_item: The inventory item to release reservation for
            quantity: Quantity to release
            
        Returns:
            bool: True if release was successful, False otherwise
        """
        if (inventory_item.reserved_quantity or 0) >= quantity:
            inventory_item.reserved_quantity = (inventory_item.reserved_quantity or 0) - quantity
            inventory_item.save(update_fields=['reserved_quantity'])
            return True
        
        return False
    
    @staticmethod
    def rent_inventory_item(
        inventory_item: InventoryItem,
        quantity: int
    ) -> bool:
        """
        Rent out a quantity of an inventory item.
        
        Args:
            inventory_item: The inventory item to rent
            quantity: Quantity to rent
            
        Returns:
            bool: True if rental was successful, False otherwise
        """
        if (inventory_item.reserved_quantity or 0) >= quantity:
            inventory_item.reserved_quantity = (inventory_item.reserved_quantity or 0) - quantity
            inventory_item.rented_quantity = (inventory_item.rented_quantity or 0) + quantity
            inventory_item.save(update_fields=['reserved_quantity', 'rented_quantity'])
            return True
        
        return False
    
    @staticmethod
    def export_inventory_items():
        """
        Export all inventory items as an Excel file.
        
        Returns:
            Dataset: Export dataset ready for response
        """
        resource = InventoryResource()
        return resource.export()
    
    @staticmethod
    def _check_inventory_number(inventory_number: str) -> bool:
        """
        Check if the inventory number is valid.

        Args:
            inventory_number: Inventory number to check

        Returns:
            bool: True if valid, False otherwise
        """
        return InventoryService.is_valid_inventory_number(inventory_number)
    
    @staticmethod
    def _create_or_skip(model, field_name: str, value: str, row_number: int, request):
        """
        Create an object or skip the row if value is missing.
        
        Args:
            model: Model class to create
            field_name: Field name for the value
            value: Value to set
            row_number: Row number for logging
            request: Django request object
            
        Returns:
            Model object or None
        """
        if value:
            obj, was_created = model.objects.get_or_create(**{field_name: value}, defaults={'description': ''})
            if was_created:
                logger.info(
                    _('%(model)s "%(value)s" created automatically.') %
                    {'model': model._meta.verbose_name, 'value': value}
                )
            return obj
        else:
            log_msg = _('%(model)s name is missing in row %(row)s. Skipping...') % {
                'model': model._meta.verbose_name, 'row': row_number
            }
            logger.warning(log_msg)
            messages.warning(request, log_msg)
            return None
    
    @staticmethod
    def validate_inventory_file(file) -> None:
        """
        Validate the inventory file and check column naming.
        
        Args:
            file: File to validate
            
        Raises:
            ValidationError: If file is invalid
        """
        wb = load_workbook(file)
        errors = []

        ws = wb.worksheets[0]
        rows = ws.rows
        header = next(rows)
        required_columns = {
            0: 'inventory_number',
            1: 'description',
            2: 'serial_number',
            3: 'manufacturer',
            4: 'location',
            5: 'quantity',
            6: 'status',
            8: 'owner',
            9: 'inventory_number_owner',
            10: 'purchase_date',
            11: 'purchase_cost',
        }

        # Check for all required columns
        for col_idx, expected_name in required_columns.items():
            actual_name = header[col_idx].value if col_idx < len(header) else None
            if actual_name != expected_name:
                errors.append(ValidationError(
                    _('Column %(col)s should be named "%(expected)s", but got "%(actual)s"') % {
                        'col': col_idx + 1,
                        'expected': expected_name,
                        'actual': actual_name or 'missing'
                    }
                ))

        if errors:
            raise ValidationError(errors)
    
    @staticmethod
    def import_inventory_data(request, file, import_obj) -> Dict[str, Union[int, str]]:
        """
        Import inventory items from an Excel file.
        
        Args:
            request: Django request object
            file: File to import
            import_obj: InventoryImport object for tracking
            
        Returns:
            Dict: Import results with created, skipped counts and error log
        """
        created_counter = 0
        skipped_counter = 0
        error_logs = []
        error_details = []

        try:
            wb = load_workbook(file)
            ws = wb.worksheets[0]
            rows = ws.rows
            headers = next(rows)

            for row in rows:
                try:
                    inventory_number = str(row[0].value or '').strip()
                    row_data = [str(cell.value) if cell.value is not None else '' for cell in row]

                    if not inventory_number or not InventoryService._check_inventory_number(inventory_number):
                        error_msg = _(f'Invalid inventory number "{inventory_number}"')
                        error_logs.append(f'Row {row[0].row}: {error_msg}')
                        error_details.append([row[0].row, error_msg] + row_data)
                        skipped_counter += 1
                        continue

                    # Get values from row
                    description = str(row[1].value or '').strip() or None
                    serial_number = str(row[2].value or '').strip() or None
                    manufacturer_name = str(row[3].value or '').strip() or None
                    location_str = str(row[4].value or '').strip()

                    # Resolve location by path in dictionary; do not create new ones here
                    def _resolve_location(path_str: str):
                        if not path_str:
                            return None
                        raw = str(path_str).replace('/', '->')
                        parts = [p.strip() for p in raw.split('->') if p and p.strip()]
                        if not parts:
                            return None
                        parent = None
                        for name in parts:
                            try:
                                node = Location.objects.get(parent=parent, name=name)
                            except Location.DoesNotExist:
                                return None
                            parent = node
                        return parent

                    loc_obj = _resolve_location(location_str)
                    if not loc_obj:
                        error_msg = _('Location "%(location)s" not found in dictionary') % {
                            'location': location_str
                        }
                        error_logs.append(f'Row {row[0].row}: {error_msg}')
                        error_details.append([row[0].row, error_msg] + row_data)
                        skipped_counter += 1
                        continue

                    # Handle quantity with default value 1
                    try:
                        quantity = int(float(str(row[5].value or '1').strip()))
                    except (ValueError, TypeError):
                        quantity = 1

                    # Convert German status to English
                    status_map = {
                        _('in Betrieb'): 'in_stock',
                        _('defekt'): 'defect',
                        _('ausgemustert'): 'written_off',
                        _('verliehen'): 'rented',
                        _('Ausleihe'): 'rented'
                    }
                    raw_status = str(row[6].value or '').strip()
                    status = status_map.get(raw_status, 'in_stock')

                    owner_name = str(row[8].value or '').strip() or None
                    inventory_number_owner = str(row[9].value or '').strip() or None

                    # Handle purchase date
                    purchase_date = None
                    if row[10].value:
                        try:
                            if isinstance(row[10].value, datetime):
                                purchase_date = row[10].value.date()
                            else:
                                purchase_date = datetime.strptime(str(row[10].value).strip(), '%Y-%m-%d').date()
                        except ValueError:
                            pass

                    # Handle cost
                    purchase_cost = None
                    if row[11].value:
                        try:
                            purchase_cost = float(str(row[11].value).strip().replace(',', '.'))
                        except ValueError:
                            pass

                    # Create or get manufacturer
                    manufacturer = None
                    if manufacturer_name:
                        manufacturer = InventoryService._create_or_skip(
                            Manufacturer, 'name', manufacturer_name, row[0].row, request
                        )
                        if not manufacturer:
                            skipped_counter += 1
                            continue

                    # Create or get owner
                    owner = None
                    if owner_name:
                        owner = InventoryService._create_or_skip(
                            Organization, 'name', owner_name, row[0].row, request
                        )
                        if not owner:
                            skipped_counter += 1
                            continue

                    # Check for duplicates
                    if InventoryItem.objects.filter(inventory_number=inventory_number).exists():
                        logger.info(
                            _('Inventory item with number "%(number)s" already exists. Skipping...') %
                            {'number': inventory_number}
                        )
                        messages.warning(
                            request,
                            _('Inventory item with number "%(number)s" already exists. Skipping...') %
                            {'number': inventory_number}
                        )
                        skipped_counter += 1
                        continue

                    # Create record
                    InventoryItem.objects.create(
                        inventory_number=inventory_number,
                        description=description,
                        serial_number=serial_number,
                        manufacturer=manufacturer,
                        location=loc_obj,
                        quantity=quantity,
                        status=status,
                        owner=owner,
                        inventory_number_owner=inventory_number_owner,
                        purchase_date=purchase_date,
                        purchase_cost=purchase_cost
                    )
                    created_counter += 1

                except Exception as e:
                    error_msg = str(e)
                    error_logs.append(_('Row %(row)s: %(error)s') % {
                        'row': row[0].row,
                        'error': error_msg
                    })
                    error_details.append([row[0].row, error_msg] + row_data)
                    skipped_counter += 1

            # If there are errors, create Excel file
            if error_details:
                error_wb = Workbook()
                ws = error_wb.active
                ws.title = _("Import Errors")

                # Headers
                ws.append([_('Row'), _('Error')] + [cell.value for cell in headers])

                # Error data
                for error_row in error_details:
                    ws.append(error_row)

                # Save to temporary file
                with NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
                    error_wb.save(tmp.name)
                    # Save to model
                    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                    filename = _('import_errors_%(timestamp)s.xlsx') % {
                        'timestamp': timestamp
                    }
                    import_obj.error_log_file.save(filename, File(open(tmp.name, 'rb')))

        except Exception as e:
            error_msg = str(e)
            error_logs.append(_('Row %(row)s: %(error)s') % {
                'row': row[0].row if 'row' in locals() else 'Unknown',
                'error': error_msg
            })

        return {
            'created': created_counter,
            'skipped': skipped_counter,
            'error_log': '\n'.join(error_logs) if error_logs else _('No errors')
        }
    
    @staticmethod
    def _read_xlsx(file) -> List[Dict[str, str]]:
        """
        Read data from an XLSX file.

        Args:
            file: XLSX file object.

        Returns:
            List of dictionaries with row data.
        """
        wb = load_workbook(file)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]  # type: ignore
        data = []
        for row in ws.iter_rows(min_row=2):  # type: ignore
            data.append({headers[i]: str(cell.value) if cell.value is not None else ''
                        for i, cell in enumerate(row)})
        return data
    
    @staticmethod
    def _parse_date(date_str: str) -> datetime.date:
        """
        Parse a date from various formats.

        Args:
            date_str: String with the date to parse.

        Returns:
            datetime.date object.

        Raises:
            ValueError: If the date format is not supported.
        """
        date_formats = [
            '%Y-%m-%d',  # 2024-03-20
            '%d.%m.%Y',  # 20.03.2024
            '%d/%m/%Y',  # 20/03/2024
            '%Y/%m/%d',  # 2024/03/20
        ]

        for date_format in date_formats:
            try:
                return datetime.strptime(date_str.strip(), date_format).date()
            except (ValueError, TypeError):
                continue

        raise ValueError(
            _('Invalid date format for "inspection_date". Supported formats: YYYY-MM-DD, DD.MM.YYYY, DD/MM/YYYY, YYYY/MM/DD. Received: "%(date_str)s"') % {'date_str': date_str}
        )
    
    @staticmethod
    def validate_inspection_file(file) -> None:
        """
        Validate the inspection file and check column names.

        Args:
            file: File object to validate.

        Raises:
            ValidationError: If the file format is invalid or required columns are missing.
        """
        file_extension = file.name.split('.')[-1].lower()
        headers: Optional[List[str]] = None

        if file_extension == 'csv':
            try:
                decoded_content = file.read().decode('utf-8')
            except UnicodeDecodeError:
                file.seek(0)
                decoded_content = file.read().decode('latin-1')
            file.seek(0)
            reader = csv.DictReader(decoded_content.splitlines())
            headers = reader.fieldnames
        elif file_extension == 'xlsx':
            wb = load_workbook(file)
            ws = wb.active
            if ws.max_row >= 1:  # type: ignore
                headers = [cell.value for cell in ws[1]]  # type: ignore
            else:
                headers = []
        else:
            raise ValidationError(_('File must be CSV or XLSX format'))

        if headers is None:
            raise ValidationError(_('Could not read headers from file.'))

        errors = []
        if not REQUIRED_INSPECTION_FIELDS.issubset(set(headers)):
            missing_columns = REQUIRED_INSPECTION_FIELDS - set(headers)
            errors.append(ValidationError(
                _('File must contain columns: %(columns)s. Missing: %(missing)s') % {
                    'columns': ', '.join(REQUIRED_INSPECTION_FIELDS),
                    'missing': ', '.join(missing_columns)
                }
            ))

        if errors:
            raise ValidationError(errors)
    
    @staticmethod
    @transaction.atomic
    def import_inspection_data(
        request: Optional[Any] = None,
        file: Optional[Any] = None,
        import_obj: Optional[Any] = None
    ) -> Dict[str, Union[int, str]]:
        """
        Import inspections from a CSV or XLSX file.

        Args:
            request: Django request object (can be None).
            file: File to import.
            import_obj: InspectionImport object for error logging.

        Returns:
            Dictionary with import results: {'created': int, 'skipped': int, 'error_log': str}.
        """
        created_counter = 0
        skipped_counter = 0
        error_logs: List[str] = []
        error_details: List[List[Any]] = []

        row_number_offset = 2  # Start from row 2 (row 1 is the header)

        try:
            if not file:
                raise ValueError(_("No file provided for import."))

            file.seek(0)
            file_extension = file.name.split('.')[-1].lower()

            rows_iterator: Union[csv.DictReader, List[Dict[str, str]]]
            header_row: Optional[List[str]] = None

            if file_extension == 'csv':
                try:
                    decoded_content = file.read().decode('utf-8')
                except UnicodeDecodeError:
                    file.seek(0)
                    decoded_content = file.read().decode('latin-1')
                file.seek(0)
                csv_reader = csv.reader(decoded_content.splitlines())
                header_list = next(csv_reader, None)
                if header_list is None:
                    raise ValueError(_("CSV file is empty or has no header row."))
                header_row = [h.strip() for h in header_list]
                rows_iterator = csv.DictReader(decoded_content.splitlines()[1:], fieldnames=header_row)
            elif file_extension == 'xlsx':
                rows_iterator = InventoryService._read_xlsx(file)
                if rows_iterator:
                    header_row = list(rows_iterator[0].keys()) if isinstance(rows_iterator[0], dict) else None
            else:
                raise ValidationError(_('Unsupported file format: %(ext)s') % {'ext': file_extension})

            if not header_row:
                raise ValueError(_("Could not determine headers from the file."))

            for i, row_data in enumerate(rows_iterator):
                current_row_num = i + row_number_offset
                try:
                    for req_field in REQUIRED_INSPECTION_FIELDS:
                        if req_field not in row_data or not row_data[req_field]:
                            raise ValueError(_("Missing required field: %(field)s") % {'field': req_field})

                    inv_num = row_data.get("inventory_number", "").strip() or None
                    item = InventoryItem.objects.filter(inventory_number=inv_num).first() if inv_num else None

                    date_str = row_data.get("inspection_date", "")
                    if not date_str:
                        raise ValueError(_("Field 'inspection_date' cannot be empty."))
                    date_ = InventoryService._parse_date(date_str)

                    inspection_number = row_data.get("inspection_number", "").strip()
                    if not inspection_number:
                        raise ValueError(_("Field 'inspection_number' cannot be empty."))

                    defaults = {
                        "inventory_item": item,
                        "manufacturer": row_data.get("manufacturer", "").strip(),
                        "device_type": row_data.get("device_type", "").strip(),
                        "room": row_data.get("room", "").strip(),
                        "inspection_date": date_,
                        "result": row_data.get("result", "").strip(),
                        "target_part": row_data.get("target_part", "device").strip() or "device",
                    }

                    obj, is_created = Inspection.objects.get_or_create(
                        inspection_number=inspection_number,
                        defaults=defaults
                    )

                    if not is_created:
                        updated_fields = []
                        for field, value in defaults.items():
                            if getattr(obj, field) != value:
                                setattr(obj, field, value)
                                updated_fields.append(field)
                        if updated_fields:
                            obj.save(update_fields=updated_fields)
                        skipped_counter += 1
                    else:
                        created_counter += 1

                except Exception as e:
                    error_msg = str(e)
                    error_logs.append(_('Row %(row_num)s: %(error)s') % {'row_num': current_row_num, 'error': error_msg})
                    error_row_values = [row_data.get(h, '') for h in header_row]
                    error_details.append([current_row_num, error_msg] + error_row_values)
                    skipped_counter += 1
                    if request:
                        messages.warning(request, _('Error processing row %(row_num)s: %(error)s') % {'row_num': current_row_num, 'error': error_msg})

            file.seek(0)

            if error_details and import_obj and header_row:
                error_wb = Workbook()
                ws = error_wb.active
                if ws:
                    ws.title = _("Import Errors")
                    ws.append([_('Row'), _('Error')] + header_row)
                    for error_row_entry in error_details:
                        ws.append(error_row_entry)
                    with NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
                        error_wb.save(tmp.name)
                        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                        filename = _('import_errors_%(timestamp)s.xlsx') % {
                            'timestamp': timestamp
                        }
                        if hasattr(import_obj, 'error_log_file'):
                            with open(tmp.name, 'rb') as tmp_file_to_save:
                                import_obj.error_log_file.save(filename, File(tmp_file_to_save))

        except Exception as e:
            logger.error(f"Inspection import failed: {str(e)}", exc_info=True)
            error_msg = str(e)
            if request:
                messages.error(request, _('Import failed: %(error)s') % {'error': error_msg})
            error_logs.append(_('Fatal Error: %(error)s') % {'error': error_msg})

        final_error_log = '\n'.join(error_logs) if error_logs else _('No errors during import.')
        if import_obj and hasattr(import_obj, 'error_log') and not getattr(import_obj, 'error_log_file', None):
            import_obj.error_log = final_error_log
            import_obj.save(update_fields=['error_log'])

        return {
            'created': created_counter,
            'skipped': skipped_counter,
            'error_log': final_error_log
        }
    
    @staticmethod
    def return_inventory_item(
        inventory_item: InventoryItem,
        quantity: int
    ) -> bool:
        """
        Return a quantity of an inventory item.
        
        Args:
            inventory_item: The inventory item to return
            quantity: Quantity to return
            
        Returns:
            bool: True if return was successful, False otherwise
        """
        if (inventory_item.rented_quantity or 0) >= quantity:
            inventory_item.rented_quantity = (inventory_item.rented_quantity or 0) - quantity
            inventory_item.save(update_fields=['rented_quantity'])
            return True
        
        return False