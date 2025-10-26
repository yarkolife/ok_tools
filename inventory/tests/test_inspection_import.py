"""
Unit tests for inventory/inspection_import.py
"""

from datetime import datetime
from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from django.forms import ValidationError
from unittest.mock import patch, MagicMock
import tempfile
import os
import io

from inventory.inspection_import import (
    _read_xlsx, validate, _parse_date, inspection_import
)
from inventory.models import Inspection, InventoryItem, InspectionImport


class ReadXlsxTest(TestCase):
    """Test cases for _read_xlsx function."""
    
    def test_read_xlsx_with_valid_file(self):
        """Test reading data from a valid XLSX file."""
        # Create a temporary XLSX file with test data
        import openpyxl
        from openpyxl import Workbook
        
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Sheet"
        
        # Add headers
        headers = ["inspection_number", "inspection_date", "inventory_number"]
        ws.append(headers)
        
        # Add data rows
        ws.append(["INS-001", "2024-01-15", "OK-001"])
        ws.append(["INS-002", "2024-02-20", "OK-002"])
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            wb.save(tmp.name)
            temp_file_path = tmp.name
        
        try:
            # Open and read the file
            with open(temp_file_path, 'rb') as file:
                file_obj = SimpleUploadedFile(
                    "test.xlsx", 
                    file.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            
            # Test the function
            result = _read_xlsx(file_obj)
            
            # Verify the result
            self.assertEqual(len(result), 2)
            self.assertEqual(result[0]["inspection_number"], "INS-001")
            self.assertEqual(result[0]["inspection_date"], "2024-01-15")
            self.assertEqual(result[0]["inventory_number"], "OK-001")
            self.assertEqual(result[1]["inspection_number"], "INS-002")
            self.assertEqual(result[1]["inspection_date"], "2024-02-20")
            self.assertEqual(result[1]["inventory_number"], "OK-002")
        finally:
            # Clean up
            os.unlink(temp_file_path)
    
    def test_read_xlsx_with_empty_file(self):
        """Test reading data from an empty XLSX file."""
        import openpyxl
        from openpyxl import Workbook
        
        # Create empty workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Sheet"
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            wb.save(tmp.name)
            temp_file_path = tmp.name
        
        try:
            # Open and read the file
            with open(temp_file_path, 'rb') as file:
                file_obj = SimpleUploadedFile(
                    "test.xlsx", 
                    file.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            
            # Test the function
            result = _read_xlsx(file_obj)
            
            # Verify the result
            self.assertEqual(len(result), 0)
        finally:
            # Clean up
            os.unlink(temp_file_path)


class ValidateTest(TestCase):
    """Test cases for validate function."""
    
    def test_validate_csv_with_required_columns(self):
        """Test validating a CSV file with required columns."""
        # Create a CSV content with required columns
        csv_content = "inspection_number,inspection_date\nINS-001,2024-01-15\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Test should not raise any exception
        try:
            validate(file_obj)
        except ValidationError:
            self.fail("validate() raised ValidationError unexpectedly for valid CSV file")
    
    def test_validate_csv_missing_required_columns(self):
        """Test validating a CSV file with missing required columns."""
        # Create a CSV content without required columns
        csv_content = "other_column,another_column\nvalue1,value2\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Test should raise ValidationError
        with self.assertRaises(ValidationError):
            validate(file_obj)
    
    def test_validate_xlsx_with_required_columns(self):
        """Test validating an XLSX file with required columns."""
        import openpyxl
        from openpyxl import Workbook
        
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Sheet"
        
        # Add required headers
        headers = ["inspection_number", "inspection_date"]
        ws.append(headers)
        
        # Add data row
        ws.append(["INS-001", "2024-01-15"])
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            wb.save(tmp.name)
            temp_file_path = tmp.name
        
        try:
            # Open and read the file
            with open(temp_file_path, 'rb') as file:
                file_obj = SimpleUploadedFile(
                    "test.xlsx", 
                    file.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            
            # Test should not raise any exception
            try:
                validate(file_obj)
            except ValidationError:
                self.fail("validate() raised ValidationError unexpectedly for valid XLSX file")
        finally:
            # Clean up
            os.unlink(temp_file_path)
    
    def test_validate_xlsx_missing_required_columns(self):
        """Test validating an XLSX file with missing required columns."""
        import openpyxl
        from openpyxl import Workbook
        
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Sheet"
        
        # Add non-required headers
        headers = ["other_column", "another_column"]
        ws.append(headers)
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            wb.save(tmp.name)
            temp_file_path = tmp.name
        
        try:
            # Open and read the file
            with open(temp_file_path, 'rb') as file:
                file_obj = SimpleUploadedFile(
                    "test.xlsx", 
                    file.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            
            # Test should raise ValidationError
            with self.assertRaises(ValidationError):
                validate(file_obj)
        finally:
            # Clean up
            os.unlink(temp_file_path)
    
    def test_validate_unsupported_file_format(self):
        """Test validating a file with unsupported format."""
        # Create a file-like object with unsupported extension
        file_obj = SimpleUploadedFile(
            "test.txt", 
            b"content",
            content_type="text/plain"
        )
        
        # Test should raise ValidationError
        with self.assertRaises(ValidationError):
            validate(file_obj)


class ParseDateTest(TestCase):
    """Test cases for _parse_date function."""
    
    def test_parse_date_yyyy_mm_dd(self):
        """Test parsing date in YYYY-MM-DD format."""
        result = _parse_date("2024-01-15")
        expected = datetime(2024, 1, 15).date()
        self.assertEqual(result, expected)
    
    def test_parse_date_dd_mm_yyyy(self):
        """Test parsing date in DD.MM.YYYY format."""
        result = _parse_date("15.01.2024")
        expected = datetime(2024, 1, 15).date()
        self.assertEqual(result, expected)
    
    def test_parse_date_dd_slash_mm_slash_yyyy(self):
        """Test parsing date in DD/MM/YYYY format."""
        result = _parse_date("15/01/2024")
        expected = datetime(2024, 1, 15).date()
        self.assertEqual(result, expected)
    
    def test_parse_date_yyyy_slash_mm_slash_dd(self):
        """Test parsing date in YYYY/MM/DD format."""
        result = _parse_date("2024/01/15")
        expected = datetime(2024, 1, 15).date()
        self.assertEqual(result, expected)
    
    def test_parse_date_invalid_format(self):
        """Test parsing date with invalid format."""
        with self.assertRaises(ValueError):
            _parse_date("invalid-date-format")
    
    def test_parse_date_empty_string(self):
        """Test parsing empty date string."""
        with self.assertRaises(ValueError):
            _parse_date("")


class InspectionImportTest(TestCase):
    """Test cases for inspection_import function."""
    
    def setUp(self):
        """Set up test data."""
        # Create a test location first
        from inventory.models import Location
        location = Location.objects.create(name='Test Location')
        
        # Create a test inventory item
        self.inventory_item = InventoryItem.objects.create(
            inventory_number='OK-001',
            description='Test Item',
            quantity=1,
            status='in_stock',
            location=location
        )
    
    def test_inspection_import_csv_success(self):
        """Test importing inspections from a valid CSV file."""
        # Create a CSV content
        csv_content = "inspection_number,inspection_date,inventory_number,manufacturer,device_type,room,result,target_part\n"
        csv_content += "INS-001,2024-01-15,OK-001,Test Manufacturer,Test Device,Room A,Passed,device\n"
        csv_content += "INS-002,15.02.2024,OK-001,Test Manufacturer 2,Test Device 2,Room B,Failed,cable\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Call the function
        result = inspection_import(request=None, file=file_obj, import_obj=None)
        
        # Verify the result
        self.assertEqual(result['created'], 2)
        self.assertEqual(result['skipped'], 0)
        # Account for localization - the message can be in English or German
        self.assertTrue('No errors during import.' in result['error_log'] or 'Keine Fehler beim Import.' in result['error_log'])
        
        # Verify that inspections were created
        inspections = Inspection.objects.order_by('inspection_number')
        self.assertEqual(inspections.count(), 2)
        
        # Check first inspection
        inspection1 = inspections[0]
        self.assertEqual(inspection1.inspection_number, "INS-001")
        self.assertEqual(inspection1.inspection_date, datetime(2024, 1, 15).date())
        self.assertEqual(inspection1.inventory_item, self.inventory_item)
        self.assertEqual(inspection1.manufacturer, "Test Manufacturer")
        self.assertEqual(inspection1.device_type, "Test Device")
        self.assertEqual(inspection1.room, "Room A")
        self.assertEqual(inspection1.result, "Passed")
        self.assertEqual(inspection1.target_part, "device")
        
        # Check second inspection
        inspection2 = inspections[1]
        self.assertEqual(inspection2.inspection_number, "INS-002")
        self.assertEqual(inspection2.inspection_date, datetime(2024, 2, 15).date())
        self.assertEqual(inspection2.inventory_item, self.inventory_item)
        self.assertEqual(inspection2.manufacturer, "Test Manufacturer 2")
        self.assertEqual(inspection2.device_type, "Test Device 2")
        self.assertEqual(inspection2.room, "Room B")
        self.assertEqual(inspection2.result, "Failed")
        self.assertEqual(inspection2.target_part, "cable")
    
    def test_inspection_import_xlsx_success(self):
        """Test importing inspections from a valid XLSX file."""
        import openpyxl
        from openpyxl import Workbook
        
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Sheet"
        
        # Add headers
        headers = ["inspection_number", "inspection_date", "inventory_number", "manufacturer", "device_type", "room", "result", "target_part"]
        ws.append(headers)
        
        # Add data rows
        ws.append(["INS-003", "2024-03-15", "OK-001", "Test Manufacturer 3", "Test Device 3", "Room C", "Passed", "psu"])
        ws.append(["INS-004", "20.04.2024", "OK-001", "Test Manufacturer 4", "Test Device 4", "Room D", "Failed", "device"])
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            wb.save(tmp.name)
            temp_file_path = tmp.name
        
        try:
            # Open and read the file
            with open(temp_file_path, 'rb') as file:
                file_obj = SimpleUploadedFile(
                    "test.xlsx", 
                    file.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            
            # Call the function
            result = inspection_import(request=None, file=file_obj, import_obj=None)
            
            # Verify the result
            self.assertEqual(result['created'], 2)
            self.assertEqual(result['skipped'], 0)
            # Account for localization - the message can be in English or German
            self.assertTrue('No errors during import.' in result['error_log'] or 'Keine Fehler beim Import.' in result['error_log'])
            
            # Verify that inspections were created
            inspections = Inspection.objects.order_by('inspection_number')
            self.assertEqual(inspections.count(), 2)
            
            # Check first inspection
            inspection1 = inspections[0]
            self.assertEqual(inspection1.inspection_number, "INS-003")
            self.assertEqual(inspection1.inspection_date, datetime(2024, 3, 15).date())
            self.assertEqual(inspection1.inventory_item, self.inventory_item)
            self.assertEqual(inspection1.manufacturer, "Test Manufacturer 3")
            self.assertEqual(inspection1.device_type, "Test Device 3")
            self.assertEqual(inspection1.room, "Room C")
            self.assertEqual(inspection1.result, "Passed")
            self.assertEqual(inspection1.target_part, "psu")
            
            # Check second inspection
            inspection2 = inspections[1]
            self.assertEqual(inspection2.inspection_number, "INS-004")
            self.assertEqual(inspection2.inspection_date, datetime(2024, 4, 20).date())
            self.assertEqual(inspection2.inventory_item, self.inventory_item)
            self.assertEqual(inspection2.manufacturer, "Test Manufacturer 4")
            self.assertEqual(inspection2.device_type, "Test Device 4")
            self.assertEqual(inspection2.room, "Room D")
            self.assertEqual(inspection2.result, "Failed")
            self.assertEqual(inspection2.target_part, "device")
        finally:
            # Clean up
            os.unlink(temp_file_path)
    
    def test_inspection_import_with_existing_inspection(self):
        """Test importing inspections when some already exist (should update)."""
        # Create an existing inspection
        existing_inspection = Inspection.objects.create(
            inspection_number="INS-001",
            inspection_date=datetime(2023, 1, 1).date(),
            manufacturer="Old Manufacturer",
            device_type="Old Device",
            room="Old Room",
            result="Old Result",
            target_part="device"
        )
        
        # Create a CSV content with the same inspection number but different data
        csv_content = "inspection_number,inspection_date,inventory_number,manufacturer,device_type,room,result,target_part\n"
        csv_content += "INS-001,2024-01-15,OK-001,New Manufacturer,New Device,New Room,New Result,cable\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Call the function
        result = inspection_import(request=None, file=file_obj, import_obj=None)
        
        # Verify the result
        self.assertEqual(result['created'], 0)  # No new inspections created
        self.assertEqual(result['skipped'], 1)  # One inspection was updated (counted as skipped)
        # Account for localization - the message can be in English or German
        self.assertTrue('No errors during import.' in result['error_log'] or 'Keine Fehler beim Import.' in result['error_log'])
        
        # Refresh the existing inspection from database
        existing_inspection.refresh_from_db()
        
        # Verify that the inspection was updated
        self.assertEqual(existing_inspection.manufacturer, "New Manufacturer")
        self.assertEqual(existing_inspection.device_type, "New Device")
        self.assertEqual(existing_inspection.room, "New Room")
        self.assertEqual(existing_inspection.result, "New Result")
        self.assertEqual(existing_inspection.target_part, "cable")
        self.assertEqual(existing_inspection.inspection_date, datetime(2024, 1, 15).date())
        self.assertEqual(existing_inspection.inventory_item, self.inventory_item)
    
    def test_inspection_import_missing_required_field(self):
        """Test importing inspections with missing required fields."""
        # Create a CSV content with missing inspection_date
        csv_content = "inspection_number,inventory_number,manufacturer,device_type,room,result,target_part\n"
        csv_content += "INS-001,OK-001,Test Manufacturer,Test Device,Room A,Passed,device\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Call the function
        result = inspection_import(request=None, file=file_obj, import_obj=None)
        
        # Verify the result
        self.assertEqual(result['created'], 0)
        self.assertEqual(result['skipped'], 1)  # One row failed to process
        # Account for localization - the message can be in English or German
        self.assertTrue('Row 2: Missing required field: inspection_date' in result['error_log'] or 'Zeile 2: Erforderliches Feld fehlt: inspection_date' in result['error_log'])
        
        # Verify that no inspections were created
        inspections = Inspection.objects.all()
        self.assertEqual(inspections.count(), 0)
    
    def test_inspection_import_invalid_date_format(self):
        """Test importing inspections with invalid date format."""
        # Create a CSV content with invalid date format
        csv_content = "inspection_number,inspection_date,inventory_number,manufacturer,device_type,room,result,target_part\n"
        csv_content += "INS-001,invalid-date,OK-001,Test Manufacturer,Test Device,Room A,Passed,device\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Call the function
        result = inspection_import(request=None, file=file_obj, import_obj=None)
        
        # Verify the result
        self.assertEqual(result['created'], 0)
        self.assertEqual(result['skipped'], 1)  # One row failed to process
        # Account for localization - the message can be in English or German
        self.assertTrue('Row 2: Invalid date format for "inspection_date"' in result['error_log'] or 'Zeile 2: Ungültiges Datumsformat für "inspection_date"' in result['error_log'])
        
        # Verify that no inspections were created
        inspections = Inspection.objects.all()
        self.assertEqual(inspections.count(), 0)
    
    def test_inspection_import_with_error_log_file(self):
        """Test importing inspections with error log file creation."""
        # Create a CSV content with some valid and some invalid rows
        csv_content = "inspection_number,inspection_date,inventory_number,manufacturer,device_type,room,result,target_part\n"
        csv_content += "INS-001,2024-01-15,OK-001,Test Manufacturer,Test Device,Room A,Passed,device\n"
        csv_content += "INS-002,invalid-date,OK-001,Test Manufacturer,Test Device,Room B,Failed,cable\n"  # Invalid date
        csv_content += "INS-003,2024-03-15,OK-001,Test Manufacturer,Test Device,Room C,Passed,psu\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Create a mock import object
        import_obj = MagicMock()
        import_obj.error_log_file = MagicMock()
        
        # Call the function
        result = inspection_import(request=None, file=file_obj, import_obj=import_obj)
        
        # Verify the result
        self.assertEqual(result['created'], 2)  # Two valid rows
        self.assertEqual(result['skipped'], 1)  # One invalid row
        # Account for localization - the message can be in English or German
        self.assertTrue('Row 3: Invalid date format for "inspection_date"' in result['error_log'] or 'Zeile 3: Ungültiges Datumsformat für "inspection_date"' in result['error_log'])
        
        # Verify that inspections were created for valid rows only
        inspections = Inspection.objects.all()
        self.assertEqual(inspections.count(), 2)
        
        # Check that error log file was created
        if hasattr(import_obj, 'error_log_file'):
            import_obj.error_log_file.save.assert_called_once()
    
    def test_inspection_import_no_file_provided(self):
        """Test importing inspections with no file provided."""
        # Call the function without file
        result = inspection_import(request=None, file=None, import_obj=None)
        
        # The function should return an error in the result rather than raising an exception
        self.assertIn("Keine Datei für den Import bereitgestellt", result['error_log'])
    
    def test_inspection_import_empty_csv(self):
        """Test importing inspections from an empty CSV file."""
        # Create an empty CSV content with just headers
        csv_content = "inspection_number,inspection_date,inventory_number,manufacturer,device_type,room,result,target_part\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Call the function
        result = inspection_import(request=None, file=file_obj, import_obj=None)
        
        # Verify the result
        self.assertEqual(result['created'], 0)
        self.assertEqual(result['skipped'], 0)
        # Account for localization - the message can be in English or German
        self.assertTrue('No errors during import.' in result['error_log'] or 'Keine Fehler beim Import.' in result['error_log'])
        
        # Verify that no inspections were created
        inspections = Inspection.objects.all()
        self.assertEqual(inspections.count(), 0)
    
    @patch('inventory.inspection_import.messages')
    def test_inspection_import_with_request(self, mock_messages):
        """Test importing inspections with a request object for error messages."""
        # Create a CSV content with invalid data
        csv_content = "inspection_number,inspection_date,inventory_number,manufacturer,device_type,room,result,target_part\n"
        csv_content += "INS-001,invalid-date,OK-001,Test Manufacturer,Test Device,Room A,Passed,device\n"
        
        # Create a file-like object
        file_obj = SimpleUploadedFile(
            "test.csv", 
            csv_content.encode('utf-8'),
            content_type="text/csv"
        )
        
        # Create a mock request object
        mock_request = MagicMock()
        
        # Call the function
        result = inspection_import(request=mock_request, file=file_obj, import_obj=None)
        
        # Verify the result
        self.assertEqual(result['created'], 0)
        self.assertEqual(result['skipped'], 1)
        # Account for localization - the message can be in English or German
        self.assertTrue('Row 2: Invalid date format for "inspection_date"' in result['error_log'] or 'Zeile 2: Ungültiges Datumsformat für "inspection_date"' in result['error_log'])
        
        # Verify that error message was added to the request
        mock_messages.warning.assert_called_once()
        
        # Verify that no inspections were created
        inspections = Inspection.objects.all()
        self.assertEqual(inspections.count(), 0)